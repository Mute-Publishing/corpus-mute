#!/usr/bin/env python3
"""Normalize Mute issue spreadsheet data into issue-level CSV outputs."""

from __future__ import annotations

import csv
import re
import urllib.error
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

SOURCE_URL_HTML_OVERRIDES: dict[int, str] = {
    34: "https://www.metamute.org/editorial/magazine/mute-vol-2-no.-3--naked-cities-struggle-global-slums.html",
}

MONTH_MAP = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}

VOL_NO_PATTERN = re.compile(
    r"^\s*Mute\s+Vol\.?\s*(?P<volume>\d+)\s*,\s*No\.?\s*(?P<number>[0-9A-Za-z.-]+)\s*(?P<tail>.*)$"
)

YEAR_MONTH_AT_END_PATTERN = re.compile(
    r"^(?P<prefix>.*?)(?P<year>\d{4})\s+(?P<month>[A-Za-z]+)\s*$"
)

ISSN_TOKEN_PATTERN = re.compile(r"^\s*(?P<issn>[0-9Xx-]+)")
ISBN_PATTERN = re.compile(r"ISBN\s*(?P<isbn>[0-9Xx-]+)", re.IGNORECASE)


@dataclass
class RowResult:
    source_row: int
    inventory_number: Optional[int]
    issue_raw: str
    source_url_raw: str
    parsed: bool
    exclusion_reason: str
    volume: Optional[int] = None
    number: str = ""
    issue_title: str = ""
    publication_year: Optional[int] = None
    publication_month: str = ""
    publication_month_num: Optional[int] = None
    publication_date_iso: str = ""
    issn_token: str = ""
    issn_print: str = ""
    issue_identifier_local: str = ""
    isbn: str = ""


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_issue_text(value: str) -> str:
    text = value.replace("\u2013", "-").replace("\u2014", "-").replace("\u2212", "-")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def parse_header_and_tail(issue_text: str) -> tuple[Optional[int], str, str]:
    match = VOL_NO_PATTERN.match(issue_text)
    if not match:
        return None, "", ""
    volume = int(match.group("volume"))
    number = match.group("number")
    tail = match.group("tail").strip()
    if tail.startswith("-"):
        tail = tail[1:].strip()
    return volume, number, tail


def parse_title_year_month(pre_issn_text: str) -> tuple[str, Optional[int], str, Optional[int]]:
    match = YEAR_MONTH_AT_END_PATTERN.match(pre_issn_text.strip())
    if not match:
        return "", None, "", None

    title_prefix = match.group("prefix").strip().strip("- ").strip()
    year = int(match.group("year"))
    month_raw = match.group("month")
    month_num = MONTH_MAP.get(month_raw.lower())
    return title_prefix, year, month_raw, month_num


def parse_issue_row(source_row: int, inventory_raw: object, issue_raw: object, source_url_raw: object) -> RowResult:
    issue = clean_text(issue_raw)
    source_url = clean_text(source_url_raw)

    inventory_number: Optional[int]
    if inventory_raw is None or clean_text(inventory_raw) == "":
        inventory_number = None
    else:
        inventory_number = int(float(inventory_raw))

    result = RowResult(
        source_row=source_row,
        inventory_number=inventory_number,
        issue_raw=issue,
        source_url_raw=source_url,
        parsed=False,
        exclusion_reason="",
    )

    if inventory_number is None:
        result.exclusion_reason = "missing_inventory_number"
        return result

    normalized = normalize_issue_text(issue)
    volume, number, tail = parse_header_and_tail(normalized)
    if volume is None:
        result.exclusion_reason = "pattern_not_matched"
        return result

    if ", ISSN" not in tail:
        result.exclusion_reason = "issn_segment_missing"
        return result

    pre_issn, post_issn = tail.split(", ISSN", 1)
    title, year, month_raw, month_num = parse_title_year_month(pre_issn)
    if year is None:
        result.exclusion_reason = "year_month_not_matched"
        return result
    if month_num is None:
        result.exclusion_reason = f"unknown_month:{month_raw}"
        return result

    issn_match = ISSN_TOKEN_PATTERN.match(post_issn)
    if not issn_match:
        result.exclusion_reason = "issn_token_not_matched"
        return result
    issn_token = issn_match.group("issn")

    isbn_match = ISBN_PATTERN.search(normalized)
    isbn = isbn_match.group("isbn") if isbn_match else ""
    issn_parts = issn_token.split("-")
    if len(issn_parts) >= 2:
        issn_print = "-".join(issn_parts[:2])
        issue_local = "-".join(issn_parts[2:])
    else:
        issn_print = issn_token
        issue_local = ""

    result.parsed = True
    result.volume = volume
    result.number = number
    result.issue_title = title
    result.publication_year = year
    result.publication_month = month_raw
    result.publication_month_num = month_num
    result.publication_date_iso = f"{year:04d}-{month_num:02d}"
    result.issn_token = issn_token
    result.issn_print = issn_print
    result.issue_identifier_local = issue_local
    result.isbn = isbn
    return result


def get_cell_hyperlink_target(cell) -> str:
    """Return hyperlink target from an openpyxl cell if present."""
    if cell.hyperlink and cell.hyperlink.target:
        return str(cell.hyperlink.target).strip()
    return clean_text(cell.value)


def check_url_status(url: str) -> str:
    """Return a compact status string for URL health checks."""
    if not url:
        return "missing"

    req = urllib.request.Request(
        url,
        headers={"User-Agent": "Mozilla/5.0 (compatible; CorpusMuteBot/1.0)"},
    )
    try:
        with urllib.request.urlopen(req, timeout=15) as response:
            code = getattr(response, "status", 0)
            if code == 200:
                return "ok"
            return f"non_200:{code}"
    except urllib.error.HTTPError as error:
        return f"http_error:{error.code}"
    except Exception as error:  # pragma: no cover - network/runtime dependent
        return f"error:{type(error).__name__}"


def write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    repo_root = Path(__file__).resolve().parents[2]
    source_path = repo_root / "mute-er-model" / "data.xlsx"
    normalized_path = repo_root / "mute-er-model" / "output" / "issues_normalized.csv"
    qa_report_path = repo_root / "mute-er-model" / "output" / "issues_qa_report.csv"

    workbook = load_workbook(source_path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    parsed_rows: list[RowResult] = []
    excluded_rows: list[RowResult] = []

    for source_row in range(2, sheet.max_row + 1):
        inventory = sheet.cell(source_row, 1).value
        issue_raw = sheet.cell(source_row, 2).value
        source_url_cell = sheet.cell(source_row, 3)
        source_url_raw = get_cell_hyperlink_target(source_url_cell)

        if clean_text(inventory) == "" and clean_text(issue_raw) == "" and source_url_raw == "":
            continue

        result = parse_issue_row(source_row, inventory, issue_raw, source_url_raw)
        if result.parsed:
            parsed_rows.append(result)
        else:
            excluded_rows.append(result)

    seen_sort_keys: set[str] = set()
    seen_identifier_keys: set[str] = set()
    normalized_records: list[dict[str, object]] = []

    for result in sorted(
        parsed_rows,
        key=lambda item: (
            item.publication_year or 0,
            item.publication_month_num or 0,
            item.volume or 0,
            float(item.number) if item.number.replace(".", "", 1).isdigit() else item.number,
            item.inventory_number or 0,
        ),
    ):
        number_value = result.number
        try:
            number_numeric = float(number_value)
        except ValueError:
            number_numeric = None

        printed_part = (
            f"v{result.volume:02d}-n{number_numeric:04.1f}" if number_numeric is not None else f"v{result.volume:02d}-n{number_value}"
        )
        sort_key_chronological = f"{result.publication_date_iso}-{result.inventory_number:03d}"
        sort_key_printed = printed_part
        sort_key_final = f"{sort_key_chronological}-{sort_key_printed}"

        duplicate_key = f"{result.volume}|{result.number}|{result.publication_date_iso}|{result.issue_title.lower()}"
        qa_possible_duplicate = duplicate_key in seen_identifier_keys
        seen_identifier_keys.add(duplicate_key)

        qa_sort_key_collision = sort_key_final in seen_sort_keys
        seen_sort_keys.add(sort_key_final)

        normalized_records.append(
            {
                "source_row": result.source_row,
                "source_inventory_number": result.inventory_number,
                "issue_label_raw": result.issue_raw,
                "source_url_raw": result.source_url_raw,
                "source_url_html": (
                    SOURCE_URL_HTML_OVERRIDES.get(result.inventory_number)
                    or (f"{result.source_url_raw}.html" if result.source_url_raw else "")
                ),
                "source_url_html_status": "",
                "series_title": "Mute",
                "issue_title": result.issue_title,
                "volume": result.volume,
                "number": result.number,
                "publication_year": result.publication_year,
                "publication_month": result.publication_month,
                "publication_month_num": result.publication_month_num,
                "publication_date_iso": result.publication_date_iso,
                "issn_print": result.issn_print,
                "issn_issue_token": result.issn_token,
                "issue_identifier_local": result.issue_identifier_local,
                "isbn": result.isbn,
                "cover_url_or_reference": "",
                "language": "en",
                "sort_key_chronological": sort_key_chronological,
                "sort_key_printed": sort_key_printed,
                "sort_key_final": sort_key_final,
                "qa_needs_title_check": "yes",
                "qa_needs_date_check": "yes",
                "qa_needs_numbering_check": "yes",
                "qa_needs_cover_check": "yes",
                "qa_needs_language_check": "yes",
                "qa_possible_duplicate": "yes" if qa_possible_duplicate else "",
                "qa_sort_key_collision": "yes" if qa_sort_key_collision else "",
                "qa_notes": "",
            }
        )

    normalized_fields = [
        "source_row",
        "source_inventory_number",
        "issue_label_raw",
        "source_url_raw",
        "source_url_html",
        "source_url_html_status",
        "series_title",
        "issue_title",
        "volume",
        "number",
        "publication_year",
        "publication_month",
        "publication_month_num",
        "publication_date_iso",
        "issn_print",
        "issn_issue_token",
        "issue_identifier_local",
        "isbn",
        "cover_url_or_reference",
        "language",
        "sort_key_chronological",
        "sort_key_printed",
        "sort_key_final",
        "qa_needs_title_check",
        "qa_needs_date_check",
        "qa_needs_numbering_check",
        "qa_needs_cover_check",
        "qa_needs_language_check",
        "qa_possible_duplicate",
        "qa_sort_key_collision",
        "qa_notes",
    ]

    for row in normalized_records:
        row["source_url_html_status"] = check_url_status(str(row["source_url_html"]))

    write_csv(normalized_path, normalized_fields, normalized_records)

    qa_rows: list[dict[str, object]] = []
    for row in excluded_rows:
        qa_rows.append(
            {
                "source_row": row.source_row,
                "source_inventory_number": "" if row.inventory_number is None else row.inventory_number,
                "issue_label_raw": row.issue_raw,
                "source_url_raw": row.source_url_raw,
                "qa_status": "excluded",
                "qa_reason": row.exclusion_reason,
                "qa_action": "review manually",
            }
        )

    # Add summary rows to QA report for quick validation.
    qa_rows.append(
        {
            "source_row": "SUMMARY",
            "source_inventory_number": "",
            "issue_label_raw": "",
            "source_url_raw": "",
            "qa_status": "normalized_count",
            "qa_reason": str(len(normalized_records)),
            "qa_action": "target is 51",
        }
    )
    qa_rows.append(
        {
            "source_row": "SUMMARY",
            "source_inventory_number": "",
            "issue_label_raw": "",
            "source_url_raw": "",
            "qa_status": "excluded_count",
            "qa_reason": str(len(excluded_rows)),
            "qa_action": "verify excluded rows",
        }
    )

    qa_fields = [
        "source_row",
        "source_inventory_number",
        "issue_label_raw",
        "source_url_raw",
        "qa_status",
        "qa_reason",
        "qa_action",
    ]

    write_csv(qa_report_path, qa_fields, qa_rows)

    print(f"Wrote {len(normalized_records)} rows to {normalized_path}")
    print(f"Wrote {len(qa_rows)} rows to {qa_report_path}")


if __name__ == "__main__":
    main()
